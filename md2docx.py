# -*- coding: utf-8 -*-
"""
md2docx.py — 把课程总结 Markdown 转成 Word 文档（python-docx）

支持：# 标题、段落、- / 1. 列表（两级）、- [ ] 待办、> 引用、| 表格 |、**加粗**、`代码`、--- 分隔线
用法：python md2docx.py summary.md ~/Desktop/couses/xxx.docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX  # noqa: F401  (保留给以后高亮用)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "PingFang SC"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)


def _set_font(style_or_run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None):
    f = style_or_run.font
    f.name = FONT
    rpr = style_or_run.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color


def _add_runs(par, text: str, bold: bool = False, color: RGBColor | None = None) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|https?://\S+)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Menlo"
        elif part.startswith("http"):
            run = par.add_run(part)
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            run.font.underline = True
        else:
            run = par.add_run(part)
        if bold:
            run.bold = True
        if color is not None and not part.startswith("http"):
            run.font.color.rgb = color


def _shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _border(par, side: str, color: str = "BFBFBF", size: int = 6) -> None:
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "4")
    el.set(qn("w:color"), color)
    bdr.append(el)
    ppr.append(bdr)


def _table(doc, rows: list[str], width_cm: float) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[1:] if not re.fullmatch(r"[\s:|-]*", "|".join(r))]
    n = len(header)
    t = doc.add_table(rows=1 + len(body), cols=n)
    t.autofit = False
    tblpr = t._tbl.tblPr
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(width_cm * 567)))
    tw.set(qn("w:type"), "dxa")
    tblpr.append(tw)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tblpr.append(borders)
    first = width_cm * 0.24 if n > 1 else width_cm
    widths = [first] + [(width_cm - first) / (n - 1)] * (n - 1) if n > 1 else [first]
    for c_idx, col in enumerate(t.columns):
        col.width = Cm(widths[c_idx])
    for r_idx, row in enumerate([header] + body):
        for c_idx in range(n):
            cell = t.cell(r_idx, c_idx)
            cell.width = Cm(widths[c_idx])
            par = cell.paragraphs[0]
            _add_runs(par, row[c_idx] if c_idx < len(row) else "", bold=(r_idx == 0))
            if r_idx == 0:
                _shade(cell, "DCE6F1")
    doc.add_paragraph()


def convert(md_path: str | Path, out_path: str | Path) -> Path:
    md_path, out_path = Path(md_path), Path(out_path).expanduser()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = sec.bottom_margin = Cm(2.1)
    text_width = 21 - 2 * 2.3

    _set_font(doc.styles["Normal"], 11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.25
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)):
        _set_font(doc.styles[name], size, bold=True, color=ACCENT if name in ("Title", "Heading 1") else RGBColor(0, 0, 0))
    for name in ("List Bullet", "List Bullet 2", "List Number"):
        _set_font(doc.styles[name], 11)

    title = md_path.stem
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.strip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            _table(doc, rows, text_width)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            if level == 1:
                title = m.group(2)
                _add_runs(doc.add_paragraph(style="Title"), m.group(2))
            else:
                _add_runs(doc.add_heading(level=level - 1), m.group(2))
        elif re.fullmatch(r"\s*-{3,}\s*", line):
            _border(doc.add_paragraph(), "bottom")
        elif (m := re.match(r"^>\s?(.*)", line)):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            _border(par, "left", "1F4E79", 12)
            _add_runs(par, m.group(1), color=GREY)
        elif (m := re.match(r"^(\s*)[-*]\s+\[( |x|X)\]\s+(.*)", line)):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            par.add_run("☑ " if m.group(2).strip() else "☐ ")
            _add_runs(par, m.group(3))
        elif (m := re.match(r"^(\s*)[-*]\s+(.*)", line)):
            style = "List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
            _add_runs(doc.add_paragraph(style=style), m.group(2))
        elif (m := re.match(r"^(\s*)\d+[.)]\s+(.*)", line)):
            style = "List Bullet 2" if len(m.group(1)) >= 2 else "List Number"
            _add_runs(doc.add_paragraph(style=style), m.group(2))
        else:
            _add_runs(doc.add_paragraph(), line.strip())
        i += 1

    doc.core_properties.title = title
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("用法：python md2docx.py <summary.md> <out.docx>")
    print(f"✅ {convert(sys.argv[1], sys.argv[2])}")
